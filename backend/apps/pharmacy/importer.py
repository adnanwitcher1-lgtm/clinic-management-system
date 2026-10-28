"""
Parse medicine lists uploaded as CSV / Excel (.xlsx) / PDF into canonical rows, then load them.

Column headers are matched loosely (e.g. "Medicine Name", "Qty", "Rate" ...), so the admin can
upload almost any reasonable sheet. PDFs are read with table extraction and, when the PDF has no
real tables, with a line-by-line text parser ("Panadol 500mg Tablet 120 2.5").
"""
import csv
import io
import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from django.db import transaction

from .models import Medicine, StockMovement, UniversalMedicine
from .services import apply_stock

ALIASES = {
    "name": {"name", "medicine", "medicine name", "drug", "drug name", "item", "item name", "product", "product name", "brand", "brand name", "description", "medicines"},
    "generic_name": {"generic", "generic name", "salt", "composition", "formula", "generic_name"},
    "strength": {"strength", "dose", "potency", "mg", "power", "dosage strength"},
    "form": {"form", "type", "dosage form", "category", "dosage", "pack type", "unit"},
    "stock_qty": {"stock", "stock qty", "qty", "quantity", "stock quantity", "available", "in stock", "stock_qty", "balance", "closing stock"},
    "price": {"price", "rate", "mrp", "unit price", "cost", "retail price", "selling price", "sale price", "amount"},
    "expiry_date": {"expiry", "expiry date", "exp", "exp date", "expiry_date", "expiration", "expire"},
    "batch_no": {"batch", "batch no", "batch number", "batch_no", "lot", "lot no"},
    "reorder_level": {"reorder", "reorder level", "min stock", "minimum stock", "reorder_level", "alert level"},
}

FORMS = {
    "tab": "Tablet", "tabs": "Tablet", "tablet": "Tablet", "tablets": "Tablet",
    "cap": "Capsule", "caps": "Capsule", "capsule": "Capsule", "capsules": "Capsule",
    "syp": "Syrup", "syrup": "Syrup", "susp": "Suspension", "suspension": "Suspension",
    "inj": "Injection", "injection": "Injection", "drops": "Drops", "drop": "Drops",
    "cream": "Cream", "oint": "Ointment", "ointment": "Ointment", "gel": "Gel", "lotion": "Lotion",
    "sachet": "Sachet", "inhaler": "Inhaler", "spray": "Spray", "powder": "Powder", "solution": "Solution",
    "suppository": "Suppository", "patch": "Patch", "infusion": "Infusion", "ors": "Sachet",
}
FORM_WORDS = "|".join(sorted(FORMS, key=len, reverse=True))
STRENGTH_RE = re.compile(r"\b\d+(?:[.,]\d+)?\s?(?:mg|mcg|µg|g|gm|ml|iu|%|mg/ml|mg/5ml|mg/\d+ml)\b(?:\s?/\s?\d+(?:\.\d+)?\s?(?:mg|ml|g))?", re.I)


def norm_header(h):
    return re.sub(r"\s+", " ", str(h or "").strip().lower().replace("_", " ").replace(".", ""))


def canonical(header):
    h = norm_header(header)
    for key, names in ALIASES.items():
        if h in {norm_header(n) for n in names}:
            return key
    return None


def norm_form(value):
    v = str(value or "").strip()
    return FORMS.get(v.lower().rstrip("."), v.title() if v else "")


def to_int(v, default=0):
    try:
        return int(float(str(v).replace(",", "").strip()))
    except (ValueError, TypeError):
        return default


def to_decimal(v, default=Decimal("0")):
    s = re.sub(r"[^0-9.\-]", "", str(v or "").replace(",", ""))
    try:
        return Decimal(s) if s else default
    except InvalidOperation:
        return default


def to_date(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v or "").strip()
    if not s:
        return None
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%m/%Y", "%b %Y", "%B %Y", "%d.%m.%Y", "%m-%Y", "%Y/%m/%d"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


# ---------------------------------------------------------------- readers
def _rows_to_dicts(rows):
    """First row that contains a recognisable 'name' column becomes the header."""
    header, out = None, []
    for row in rows:
        cells = ["" if c is None else str(c).strip() for c in row]
        if not any(cells):
            continue
        if header is None:
            mapped = [canonical(c) for c in cells]
            if "name" in mapped:
                header = mapped
            continue
        rec = {}
        for key, val in zip(header, cells):
            if key and val != "":
                rec.setdefault(key, val)
        if rec.get("name"):
            out.append(rec)
    return out


def _read_csv(data):
    text = None
    for enc in ("utf-8-sig", "utf-16", "cp1252", "latin-1"):
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    return _rows_to_dicts(csv.reader(io.StringIO(text), dialect))


def _read_xlsx(data):
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    out = []
    for ws in wb.worksheets:
        out.extend(_rows_to_dicts(ws.iter_rows(values_only=True)))
    return out


def parse_text_line(line):
    """Best-effort parse of one text line like 'Panadol 500mg Tablet 120 2.5'."""
    line = re.sub(r"\s+", " ", line).strip()
    if len(line) < 3 or not re.search(r"[A-Za-z]{3}", line):
        return None
    rec = {}
    nums = re.findall(r"(?<![\w.])\d+(?:\.\d+)?(?![\w%])", STRENGTH_RE.sub(" ", line))
    s = STRENGTH_RE.search(line)
    if s:
        rec["strength"] = s.group(0).replace(" ", "")
    f = re.search(rf"\b({FORM_WORDS})\b\.?", line, re.I)
    if f:
        rec["form"] = norm_form(f.group(1))
    name_end = min([m.start() for m in (s, f) if m] or [len(line)])
    name = line[:name_end].strip(" -–,|")
    if not name:
        # fall back to leading words before first number
        name = re.split(r"\s\d", line)[0].strip()
    name = re.sub(r"^\d+[.)]?\s+", "", name)  # leading serial number
    if not name:
        return None
    rec["name"] = name
    if len(nums) >= 2:
        rec["stock_qty"], rec["price"] = nums[-2], nums[-1]
    elif len(nums) == 1:
        rec["stock_qty"] = nums[0]
    return rec


def _read_pdf(data):
    import pdfplumber

    out, header = [], None
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        text_lines = []
        for page in pdf.pages:
            tables = page.extract_tables() or []
            for table in tables:
                rows = table
                # Continuation pages often repeat no header – reuse the one we found earlier.
                if header and not any(canonical(c) == "name" for c in (rows[0] or [])):
                    rows = [header] + rows
                got = _rows_to_dicts(rows)
                if got:
                    out.extend(got)
                    first = next((r for r in rows if r and any(canonical(c) for c in r)), None)
                    header = first or header
            if not tables:
                text_lines.extend((page.extract_text() or "").splitlines())
    if not out:
        for line in text_lines:
            rec = parse_text_line(line)
            if rec:
                out.append(rec)
    return out


def parse_upload(upload):
    """Return list of canonical dict rows from an uploaded file object."""
    name = (getattr(upload, "name", "") or "").lower()
    data = upload.read()
    if name.endswith(".pdf"):
        return _read_pdf(data)
    if name.endswith((".xlsx", ".xlsm")):
        return _read_xlsx(data)
    if name.endswith(".xls"):
        raise ValueError("Old .xls files are not supported – please save as .xlsx or .csv")
    return _read_csv(data)


# ---------------------------------------------------------------- loaders
@transaction.atomic
def import_store(rows, user=None, mode="add"):
    """
    mode="add"     -> quantities are added to existing stock (new purchase / delivery)
    mode="replace" -> quantity in the file becomes the stock level (stock-take)
    """
    created = updated = skipped = 0
    errors = []
    for i, r in enumerate(rows, start=1):
        name = (r.get("name") or "").strip()
        if not name:
            skipped += 1
            continue
        try:
            strength = (r.get("strength") or "").strip()
            form = norm_form(r.get("form"))
            qty = to_int(r.get("stock_qty"))
            med = Medicine.objects.filter(name__iexact=name, strength__iexact=strength, form__iexact=form).first()
            fields = {}
            if "price" in r:
                fields["price"] = to_decimal(r["price"])
            if r.get("expiry_date"):
                fields["expiry_date"] = to_date(r["expiry_date"])
            if r.get("batch_no"):
                fields["batch_no"] = r["batch_no"][:60]
            if r.get("generic_name"):
                fields["generic_name"] = r["generic_name"][:200]
            if r.get("reorder_level"):
                fields["reorder_level"] = to_int(r["reorder_level"], 10)
            if med is None:
                med = Medicine.objects.create(name=name[:200], strength=strength[:60], form=form[:60], stock_qty=0, **fields)
                created += 1
                delta = qty
            else:
                for k, v in fields.items():
                    setattr(med, k, v)
                med.is_active = True
                med.save()
                updated += 1
                delta = qty if mode == "add" else qty - med.stock_qty
                if "stock_qty" not in r:
                    delta = 0
            if delta:
                apply_stock(med.pk, delta, StockMovement.Reason.IMPORT, user=user, note=f"File import ({mode})", allow_negative=True)
        except Exception as exc:  # keep going, report row
            errors.append(f"Row {i} ({name}): {exc}")
    return {"created": created, "updated": updated, "skipped": skipped, "errors": errors[:25], "total_rows": len(rows)}


def import_universal(rows, batch=1000):
    before = UniversalMedicine.objects.count()
    seen, objs = set(), []
    for r in rows:
        name = (r.get("name") or "").strip()
        if not name:
            continue
        key = (name[:200], (r.get("strength") or "").strip()[:60], norm_form(r.get("form"))[:60])
        if key in seen:
            continue
        seen.add(key)
        objs.append(UniversalMedicine(name=key[0], strength=key[1], form=key[2], generic_name=(r.get("generic_name") or "").strip()[:200]))
    for i in range(0, len(objs), batch):
        UniversalMedicine.objects.bulk_create(objs[i : i + batch], ignore_conflicts=True)
    added = UniversalMedicine.objects.count() - before
    return {"created": added, "updated": 0, "skipped": len(rows) - added, "errors": [], "total_rows": len(rows)}
