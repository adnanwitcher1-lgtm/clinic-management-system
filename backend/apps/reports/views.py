import io
from collections import Counter
from datetime import timedelta
from decimal import Decimal

from django.db.models import Count, F, Q, Sum
from django.db.models.functions import TruncDate, TruncMonth
from django.http import HttpResponse
from django.utils import timezone
from rest_framework.response import Response
from rest_framework.views import APIView

from apps.billing.models import Payment
from apps.clinic.models import Patient, PrescriptionItem, Visit
from apps.core.models import ClinicSettings
from apps.pharmacy.models import Medicine, StockMovement


def is_admin(request):
    return request.user.role == "admin"


class DashboardView(APIView):
    def get(self, request):
        if not is_admin(request):
            return Response(status=403)
        today = timezone.localdate()
        cfg = ClinicSettings.load()
        visits_today = Visit.objects.filter(visit_date=today).exclude(status="cancelled")
        revenue_today = Payment.objects.filter(created_at__date=today).aggregate(t=Sum("amount"))["t"] or 0
        pending = visits_today.exclude(status="completed").count()
        low = Medicine.objects.filter(is_active=True, stock_qty__lte=F("reorder_level")).count()
        expiring = Medicine.objects.filter(is_active=True, expiry_date__isnull=False, expiry_date__lte=today + timedelta(days=cfg.expiry_alert_days)).count()

        start = today - timedelta(days=29)
        daily = {r["d"]: r["c"] for r in Visit.objects.filter(visit_date__gte=start).exclude(status="cancelled").annotate(d=TruncDate("created_at")).values("d").annotate(c=Count("id"))}
        # Visit.visit_date is a DateField: use it directly for accuracy
        daily = {r["visit_date"]: r["c"] for r in Visit.objects.filter(visit_date__gte=start).exclude(status="cancelled").values("visit_date").annotate(c=Count("id"))}
        daily_patients = [{"date": (start + timedelta(days=i)).isoformat(), "patients": daily.get(start + timedelta(days=i), 0)} for i in range(30)]

        first_month = (today.replace(day=1) - timedelta(days=330)).replace(day=1)
        rev = {r["m"].date() if hasattr(r["m"], "date") else r["m"]: r["t"] for r in Payment.objects.filter(created_at__date__gte=first_month).annotate(m=TruncMonth("created_at")).values("m").annotate(t=Sum("amount"))}
        monthly = []
        cur = today.replace(day=1)
        months = []
        for _ in range(12):
            months.append(cur)
            cur = (cur - timedelta(days=1)).replace(day=1)
        for m in reversed(months):
            monthly.append({"month": m.strftime("%b %y"), "revenue": float(rev.get(m, 0) or 0)})

        top_meds = list(
            PrescriptionItem.objects.values("name").annotate(count=Count("id")).order_by("-count")[:8]
        )
        dx = Counter()
        for d in Visit.objects.exclude(diagnosis="").order_by("-id").values_list("diagnosis", flat=True)[:5000]:
            dx[d.strip().lower().title()[:40]] += 1
        queue = {r["status"]: r["c"] for r in visits_today.values("status").annotate(c=Count("id"))}
        return Response({
            "cards": {
                "today_patients": visits_today.count(), "today_revenue": revenue_today, "pending_patients": pending,
                "low_stock": low, "expiring_soon": expiring, "total_patients": Patient.objects.count(),
                "total_visits": Visit.objects.count(),
            },
            "daily_patients": daily_patients, "monthly_revenue": monthly, "top_medicines": top_meds,
            "common_diagnoses": [{"name": k, "count": v} for k, v in dx.most_common(8)], "queue": queue,
            "currency": cfg.currency,
        })


# ---------------------------------------------------------------- records + export
def _range(p, field):
    f = {}
    if p.get("date_from"):
        f[f"{field}__gte"] = p["date_from"]
    if p.get("date_to"):
        f[f"{field}__lte"] = p["date_to"]
    return f


def dataset(kind, p):
    """Return (title, columns, rows) for the Records tabs & exports."""
    if kind == "patients":
        qs = Patient.objects.annotate(v=Count("visits")).filter(**_range(p, "created_at__date")).order_by("-id")
        cols = ["Patient No", "Name", "Age", "Gender", "Phone", "Blood group", "Allergies", "Visits", "Registered"]
        rows = [[x.patient_no, x.name, x.age, x.get_gender_display(), x.phone, x.blood_group, x.allergies, x.v, x.created_at.astimezone().strftime("%Y-%m-%d")] for x in qs[:3000]]
        return "Patients", cols, rows
    if kind == "visits":
        qs = Visit.objects.filter(**_range(p, "visit_date")).select_related("patient", "doctor").prefetch_related("payments", "prescription__items", "prescription__injections__components").order_by("-visit_date", "-token_no")
        cols = ["Date", "Token", "Patient No", "Patient", "Doctor", "Status", "Diagnosis", "Fee", "Medicines", "Injections", "Total", "Paid"]
        rows = [[str(v.visit_date), v.token_no, v.patient.patient_no, v.patient.name, v.doctor.display_name, v.get_status_display(), v.diagnosis[:80], float(v.consultation_fee), float(v.medicine_total), float(v.injection_total), float(v.total_due), float(v.paid_amount)] for v in qs[:3000]]
        return "Visits", cols, rows
    if kind == "inventory":
        qs = Medicine.objects.order_by("name")
        cols = ["Medicine", "Strength", "Form", "Stock", "Reorder level", "Price", "Stock value", "Expiry", "Batch", "Active"]
        rows = [[m.name, m.strength, m.form, m.stock_qty, m.reorder_level, float(m.price), float(m.price * max(m.stock_qty, 0)), str(m.expiry_date or ""), m.batch_no, "Yes" if m.is_active else "No"] for m in qs[:5000]]
        return "Medicine inventory", cols, rows
    if kind == "dispensing":
        qs = StockMovement.objects.filter(reason="dispense").filter(**_range(p, "created_at__date")).select_related("medicine", "visit__patient", "created_by")
        cols = ["Date/Time", "Token", "Patient", "Medicine", "Qty", "Unit price", "Amount", "Dispensed by"]
        rows = [[m.created_at.astimezone().strftime("%Y-%m-%d %H:%M"), m.visit.token_no if m.visit else "", m.visit.patient.name if m.visit else "", m.medicine.label, -m.qty, float(m.unit_price), float(m.unit_price * -m.qty), m.created_by.username if m.created_by else ""] for m in qs[:3000]]
        return "Medicine dispensing history", cols, rows
    if kind == "payments":
        qs = Payment.objects.filter(**_range(p, "created_at__date")).select_related("visit__patient", "received_by")
        cols = ["Date/Time", "Token", "Patient No", "Patient", "Amount", "Method", "Received by", "Note"]
        rows = [[x.created_at.astimezone().strftime("%Y-%m-%d %H:%M"), x.visit.token_no, x.visit.patient.patient_no, x.visit.patient.name, float(x.amount), x.get_method_display(), x.received_by.username if x.received_by else "", x.note] for x in qs[:3000]]
        return "Revenue / Payments", cols, rows
    raise ValueError("Unknown dataset")


def filtered(kind, p):
    title, cols, rows = dataset(kind, p)
    q = (p.get("q") or "").lower().strip()
    if q:
        rows = [r for r in rows if any(q in str(c).lower() for c in r)]
    return title, cols, rows


class RecordsView(APIView):
    def get(self, request):
        if not is_admin(request):
            return Response(status=403)
        try:
            title, cols, rows = filtered(request.query_params.get("type", "patients"), request.query_params)
        except ValueError:
            return Response({"detail": "Unknown record type"}, status=400)
        return Response({"title": title, "columns": cols, "rows": rows[:1000], "count": len(rows)})


def to_xlsx(title, cols, rows):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = title.replace("/", "-")[:30]
    ws.append(cols)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="3B3FC7")
        c.alignment = Alignment(horizontal="center")
    for r in rows:
        ws.append(r)
    for i, col in enumerate(cols, 1):
        width = max([len(str(col))] + [len(str(r[i - 1])) for r in rows[:200]]) + 3
        ws.column_dimensions[get_column_letter(i)].width = min(width, 45)
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def to_pdf(title, cols, rows, period=""):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    cfg = ClinicSettings.load()
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=10 * 28.35 / 10, rightMargin=28, topMargin=28, bottomMargin=28, title=title)
    st = ParagraphStyle("c", fontName="Helvetica", fontSize=7, leading=8.5)
    hd = ParagraphStyle("h", fontName="Helvetica-Bold", fontSize=7.5, textColor=colors.white)
    data = [[Paragraph(str(c), hd) for c in cols]] + [[Paragraph(str(x if x is not None else "")[:60], st) for x in r] for r in rows[:1500]]
    t = Table(data, repeatRows=1)
    t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#3b3fc7")), ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f1efff")]), ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#dcd9f5")), ("VALIGN", (0, 0), (-1, -1), "MIDDLE")]))
    head = ParagraphStyle("t", fontName="Helvetica-Bold", fontSize=14, textColor=colors.HexColor("#3b3fc7"))
    sub = ParagraphStyle("s", fontName="Helvetica", fontSize=8, textColor=colors.grey)
    doc.build([Paragraph(f"{cfg.clinic_name} — {title}", head), Paragraph(f"{period}  ·  {len(rows)} rows  ·  generated {timezone.localtime():%d %b %Y %H:%M}", sub), Spacer(1, 8), t])
    return buf.getvalue()


class ExportView(APIView):
    def get(self, request):
        if not is_admin(request):
            return Response(status=403)
        p = request.query_params
        try:
            title, cols, rows = filtered(p.get("type", "patients"), p)
        except ValueError:
            return Response({"detail": "Unknown record type"}, status=400)
        fmt = p.get("fmt", "xlsx")
        stamp = timezone.localdate().isoformat()
        name = f"{p.get('type', 'report')}-{stamp}"
        if fmt == "pdf":
            period = f"{p.get('date_from', 'start')} to {p.get('date_to', 'today')}"
            resp = HttpResponse(to_pdf(title, cols, rows, period), content_type="application/pdf")
            resp["Content-Disposition"] = f'attachment; filename="{name}.pdf"'
        else:
            resp = HttpResponse(to_xlsx(title, cols, rows), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
            resp["Content-Disposition"] = f'attachment; filename="{name}.xlsx"'
        return resp
